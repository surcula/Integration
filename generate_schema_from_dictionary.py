import re
from pathlib import Path

import openpyxl


PROJECT_DIR = Path(__file__).resolve().parent
DICTIONARY = PROJECT_DIR / "DICTIONNAIRE_DE_DONNEES.xlsx"
OUTPUT = PROJECT_DIR / "database_schema_from_dictionary.sql"

VALID_TYPES = {
    "INT",
    "VARCHAR",
    "DATE",
    "TEXT",
    "BOOL",
    "DATETIME",
    "TIME",
    "DECIMAL",
    "FLOAT",
    "DOUBLE",
}


def is_checked(value):
    return value == "☑"


def quote_name(name):
    return "`" + name.replace("`", "``") + "`"


def is_document_related(value):
    return value is not None and "document" in str(value).lower()


def clean_table(name):
    name = str(name).strip()
    name = re.sub(r"(?<!^)(?=[A-Z])", "_", name).lower()
    name = re.sub(r"_+", "_", name)
    fixes = {
        "trainings_center": "training_centers",
        "trainings_centers": "training_centers",
        "trainings_careers": "training_centers",
        "employee_training": "employee_trainings",
    }
    return fixes.get(name, name)


def clean_column(raw_name, is_pk, is_fk):
    name = str(raw_name).strip()
    name = re.sub(r"(?<!^)(?=[A-Z])", "_", name).lower()
    name = re.sub(r"_+", "_", name)
    name = name.replace("addresse", "address")

    if is_pk:
        return "id"

    if name.startswith("f_k_"):
        name = name[4:]
    if name.startswith("fk_"):
        name = name[3:]

    aliases = {
        "employee": "employee_id",
        "superior": "superior_id",
        "function": "function_id",
        "criterion": "criterion_id",
        "criteria": "criterion_id",
        "evaluation": "evaluation_id",
        "planning": "planning_id",
        "pers_evaluator": "evaluator_id",
        "city": "city_id",
        "category": "category_id",
    }
    name = aliases.get(name, name)

    if is_fk and not name.endswith("_id"):
        name += "_id"

    return name


def sql_type(column):
    column_type = column["type"]
    if column_type == "BOOL":
        return "TINYINT(1)"
    if column_type == "VARCHAR":
        length = int(column["length"]) if column["length"] else 255
        return f"VARCHAR({length})"
    if column_type == "DECIMAL":
        return "DECIMAL(10,2)"
    return column_type


def sql_default(column):
    value = column["default"]
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (int, float)):
        return f" DEFAULT {value}"

    value = str(value).strip()
    if value.upper() in {"CURRENT_TIMESTAMP", "NULL"}:
        return f" DEFAULT {value.upper()}"
    return " DEFAULT '" + value.replace("'", "''") + "'"


def parse_dictionary():
    workbook = openpyxl.load_workbook(DICTIONARY, data_only=True)
    worksheet = workbook.active
    tables = []
    current = None

    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row, 1).value
        raw_type = worksheet.cell(row, 2).value

        if not name and not raw_type:
            continue

        if name and raw_type is None:
            if is_document_related(name):
                current = None
                continue
            current = {"source": str(name).strip(), "name": clean_table(name), "columns": []}
            tables.append(current)
            continue

        if not current or not name:
            continue

        column_type = str(raw_type).strip().upper()
        if column_type not in VALID_TYPES:
            continue

        description = worksheet.cell(row, 4).value
        if is_document_related(name) or is_document_related(description):
            continue

        is_pk = is_checked(worksheet.cell(row, 7).value)
        is_fk = is_checked(worksheet.cell(row, 10).value)
        column = {
            "source": str(name).strip(),
            "type": column_type,
            "length": worksheet.cell(row, 3).value,
            "description": "" if description is None else str(description).strip(),
            "not_null": is_checked(worksheet.cell(row, 5).value),
            "default": worksheet.cell(row, 6).value,
            "pk": is_pk,
            "unique": is_checked(worksheet.cell(row, 8).value),
            "auto_increment": is_checked(worksheet.cell(row, 9).value),
            "fk": is_fk,
        }
        column["name"] = clean_column(column["source"], is_pk, is_fk)
        current["columns"].append(column)

    cleaned_tables = []
    seen_tables = set()
    for table in tables:
        if table["name"] in seen_tables:
            continue
        seen_tables.add(table["name"])

        columns = []
        seen_columns = set()
        for column in table["columns"]:
            if column["name"] in seen_columns:
                continue
            seen_columns.add(column["name"])
            columns.append(column)
        table["columns"] = columns
        cleaned_tables.append(table)

    return cleaned_tables


def apply_manual_schema_overrides(tables):
    columns_by_table = {
        table["name"]: {column["name"] for column in table["columns"]}
        for table in tables
    }

    manual_columns = {
        "evaluations": {
            "source": "status",
            "type": "VARCHAR",
            "length": 50,
            "description": "Evaluation workflow status",
            "not_null": True,
            "default": "CREATED",
            "pk": False,
            "unique": False,
            "auto_increment": False,
            "fk": False,
            "name": "status",
        },
        "evaluations_objectives": {
            "source": "status",
            "type": "VARCHAR",
            "length": 50,
            "description": "Evaluation objective workflow status",
            "not_null": True,
            "default": "TO_DO",
            "pk": False,
            "unique": False,
            "auto_increment": False,
            "fk": False,
            "name": "status",
        },
        "job_offers": [
            {
                "source": "email",
                "type": "VARCHAR",
                "length": 150,
                "description": "Job offer contact email",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "email",
            },
            {
                "source": "contact",
                "type": "VARCHAR",
                "length": 150,
                "description": "Job offer contact person",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "contact",
            },
            {
                "source": "duration",
                "type": "VARCHAR",
                "length": 100,
                "description": "Job offer duration",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "duration",
            },
            {
                "source": "number_of_open_positions",
                "type": "INT",
                "length": None,
                "description": "Number of open positions",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "number_of_open_positions",
            },
            {
                "source": "profil",
                "type": "TEXT",
                "length": None,
                "description": "Expected candidate profile",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "profil",
            },
            {
                "source": "job_description",
                "type": "TEXT",
                "length": None,
                "description": "Job description",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "job_description",
            },
            {
                "source": "requirements",
                "type": "TEXT",
                "length": None,
                "description": "Job requirements",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "requirements",
            },
            {
                "source": "comments",
                "type": "TEXT",
                "length": None,
                "description": "Internal comments",
                "not_null": False,
                "default": None,
                "pk": False,
                "unique": False,
                "auto_increment": False,
                "fk": False,
                "name": "comments",
            },
        ],
    }

    for table in tables:
        column = manual_columns.get(table["name"])
        if isinstance(column, list):
            for item in column:
                if item["name"] not in columns_by_table[table["name"]]:
                    table["columns"].append(item)
        elif column and column["name"] not in columns_by_table[table["name"]]:
            table["columns"].append(column)


def build_foreign_key_reference(table_names, pk_by_table, column):
    source = column["source"].lower().replace("addresse", "address")
    description = column["description"].lower().replace("addresse", "address")

    candidates = []
    match = re.search(r"pk\s+([a-zA-Z_]+)\s+table", description)
    if match:
        candidates.append(match.group(1).lower())

    if source.startswith("fk_"):
        source = source[3:]
    if source.endswith("_id"):
        source = source[:-3]
    candidates.append(source)

    if column["name"].endswith("_id"):
        candidates.append(column["name"][:-3])

    aliases = {
        "address": "addresses",
        "role": "roles",
        "authorization": "authorizations",
        "city": "cities",
        "department": "departments",
        "department_head": "department_heads",
        "superior": "superiors",
        "employee": "employees",
        "pers_evaluator": "employees",
        "evaluator": "employees",
        "function": "functions",
        "training": "trainings",
        "training_center": "training_centers",
        "category": "categories",
        "criterion": "criteria",
        "criteria": "criteria",
        "function_criterion": "criteria_function",
        "objective": "objectives",
        "evaluation": "evaluations",
        "planning": "plannings",
        "job_offer": "job_offers",
        "job_offers": "job_offers",
        "candidate": "candidates",
    }

    for candidate in candidates:
        reference_table = aliases.get(candidate, candidate)
        if reference_table in table_names:
            return reference_table, pk_by_table[reference_table]
    return None


def generate_sql(tables):
    table_names = {table["name"] for table in tables}
    pk_by_table = {
        table["name"]: next((column["name"] for column in table["columns"] if column["pk"]), "id")
        for table in tables
    }

    lines = [
        "CREATE DATABASE IF NOT EXISTS `erp_projet_integration`",
        "    CHARACTER SET utf8mb4",
        "    COLLATE utf8mb4_unicode_ci;",
        "",
        "USE `erp_projet_integration`;",
        "",
        "SET FOREIGN_KEY_CHECKS = 0;",
        "",
    ]

    for table in reversed(tables):
        lines.append(f"DROP TABLE IF EXISTS {quote_name(table['name'])};")

    lines.extend(["", "SET FOREIGN_KEY_CHECKS = 1;", ""])

    foreign_keys = []
    constraint_names = set()
    for table in tables:
        lines.append(f"CREATE TABLE {quote_name(table['name'])} (")
        definitions = []

        for column in table["columns"]:
            parts = [f"    {quote_name(column['name'])}", sql_type(column)]
            if column["pk"] or column["not_null"]:
                parts.append("NOT NULL")
            default_value = sql_default(column)
            if default_value:
                parts.append(default_value.strip())
            if column["auto_increment"]:
                parts.append("AUTO_INCREMENT")
            if column["pk"]:
                parts.append("PRIMARY KEY")
            definitions.append(" ".join(parts))

            if column["unique"] and not column["pk"]:
                definitions.append(
                    f"    UNIQUE KEY {quote_name('uq_' + table['name'] + '_' + column['name'])} "
                    f"({quote_name(column['name'])})"
                )

            if column["fk"]:
                reference = build_foreign_key_reference(table_names, pk_by_table, column)
                if reference:
                    base_constraint_name = ("fk_" + table["name"] + "_" + column["name"])[:58]
                    constraint_name = base_constraint_name
                    index = 2
                    while constraint_name in constraint_names:
                        suffix = "_" + str(index)
                        constraint_name = base_constraint_name[: 64 - len(suffix)] + suffix
                        index += 1
                    constraint_names.add(constraint_name)
                    foreign_keys.append((table["name"], constraint_name, column["name"], reference[0], reference[1]))

        lines.append(",\n".join(definitions))
        lines.append(") ENGINE=InnoDB;")
        lines.append("")

    for table, constraint_name, column, reference_table, reference_column in foreign_keys:
        lines.append(f"ALTER TABLE {quote_name(table)}")
        lines.append(
            f"    ADD CONSTRAINT {quote_name(constraint_name)} "
            f"FOREIGN KEY ({quote_name(column)}) "
            f"REFERENCES {quote_name(reference_table)} ({quote_name(reference_column)});"
        )
        lines.append("")

    lines.extend(
        [
            "INSERT INTO `roles` (`role_name`, `is_active`) VALUES",
            "('ADMIN', 1),",
            "('HR', 1),",
            "('SERVICE_MANAGER', 1),",
            "('SECRETARY', 1),",
            "('STAGIAIRE', 1),",
            "('STAGIAIRE 2', 1),",
            "('EMPLOYEE', 1);",
            "",
            "INSERT INTO `authorizations` (`authorization_name`, `is_active`) VALUES",
            "('hub:access', 1),",
            "('employee:create', 1),",
            "('employee:edit', 1),",
            "('employee:delete', 1),",
            "('job-offer:read', 1),",
            "('job-offer:create', 1),",
            "('job-offer:edit', 1),",
            "('job-offer:publish', 1),",
            "('job-offer:archive', 1),",
            "('job-offer:delete', 1),",
            "('candidate:read', 1),",
            "('candidate:create', 1),",
            "('candidate:edit', 1),",
            "('candidate:delete', 1),",
            "('evaluation:read', 1),",
            "('evaluation:create', 1),",
            "('evaluation:edit', 1),",
            "('evaluation:delete', 1);",
            "",
            "INSERT INTO `roles_authorization` (`role_authorization_name`, `is_active`, `role_id`, `authorization_id`)",
            "SELECT CONCAT(r.`role_name`, '_', a.`authorization_name`), 1, r.`id`, a.`id`",
            "FROM `roles` r",
            "JOIN `authorizations` a",
            "WHERE r.`role_name` = 'ADMIN';",
            "",
            "INSERT INTO `roles_authorization` (`role_authorization_name`, `is_active`, `role_id`, `authorization_id`)",
            "SELECT CONCAT(r.`role_name`, '_', a.`authorization_name`), 1, r.`id`, a.`id`",
            "FROM `roles` r",
            "JOIN `authorizations` a",
            "WHERE r.`role_name` = 'HR'",
            "  AND a.`authorization_name` IN (",
            "      'hub:access',",
            "      'employee:create', 'employee:edit', 'employee:delete',",
            "      'job-offer:read', 'job-offer:create', 'job-offer:edit',",
            "      'job-offer:publish', 'job-offer:archive', 'job-offer:delete',",
            "      'candidate:read', 'candidate:create', 'candidate:edit', 'candidate:delete',",
            "      'evaluation:read', 'evaluation:create', 'evaluation:edit', 'evaluation:delete'",
            "  );",
            "",
            "INSERT INTO `roles_authorization` (`role_authorization_name`, `is_active`, `role_id`, `authorization_id`)",
            "SELECT CONCAT(r.`role_name`, '_', a.`authorization_name`), 1, r.`id`, a.`id`",
            "FROM `roles` r",
            "JOIN `authorizations` a",
            "WHERE r.`role_name` = 'SERVICE_MANAGER'",
            "  AND a.`authorization_name` IN (",
            "      'hub:access',",
            "      'evaluation:read', 'evaluation:create', 'evaluation:edit', 'evaluation:delete'",
            "  );",
            "",
            "INSERT INTO `roles_authorization` (`role_authorization_name`, `is_active`, `role_id`, `authorization_id`)",
            "SELECT CONCAT(r.`role_name`, '_', a.`authorization_name`), 1, r.`id`, a.`id`",
            "FROM `roles` r",
            "JOIN `authorizations` a",
            "WHERE r.`role_name` IN ('EMPLOYEE', 'SECRETARY', 'STAGIAIRE', 'STAGIAIRE 2')",
            "  AND a.`authorization_name` IN ('hub:access', 'evaluation:read');",
            "",
        ]
    )

    return "\n".join(lines)


def main():
    tables = parse_dictionary()
    apply_manual_schema_overrides(tables)
    OUTPUT.write_text(generate_sql(tables), encoding="utf-8")
    print(f"Created {OUTPUT}")
    print(f"Tables: {len(tables)}")


if __name__ == "__main__":
    main()
